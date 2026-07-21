import os
from openpyxl import load_workbook
from PIL import Image
from io import BytesIO
import pandas as pd

# ==========================
# CONFIG
# ==========================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Look for the template next to this script; override with EXCEL_PATH env var.
excel_path = os.environ.get("EXCEL_PATH", os.path.join(BASE_DIR, "Shopping list.xlsx"))

BASE_OUTPUT = os.path.join(BASE_DIR, "images")
os.makedirs(BASE_OUTPUT, exist_ok=True)

# ==========================
# LOAD WORKBOOK
# ==========================
wb = load_workbook(excel_path)

print("Available sheets:", wb.sheetnames)

# ==========================
# HELPER: encontrar nombre real
# ==========================
def find_sheet(name):
    for s in wb.sheetnames:
        if name.lower() in s.lower():
            return s
    return None

# ==========================
# MAIN FUNCTION
# ==========================
def process_sheet(sheet_keyword, folder_name):

    real_name = find_sheet(sheet_keyword)

    if not real_name:
        print(f"⚠️ Sheet not found: {sheet_keyword}")
        return pd.DataFrame()

    print(f"✅ Processing: {real_name}")

    ws = wb[real_name]

    folder = os.path.join(BASE_OUTPUT, folder_name)
    os.makedirs(folder, exist_ok=True)

    image_map = {}

    # ==========================
    # EXTRACT IMAGES ✅ FIX RGB
    # ==========================
    for img in ws._images:
        try:
            row_excel = img.anchor._from.row + 1

            img_bytes = img._data()
            pil_img = Image.open(BytesIO(img_bytes))

            # ✅ FIX: convertir a RGB
            if pil_img.mode in ("RGBA", "P"):
                pil_img = pil_img.convert("RGB")

            filename = f"{row_excel}.jpg"
            path = os.path.join(folder, filename)

            pil_img.save(path)
            image_map[row_excel] = path

        except Exception as e:
            print(f"⚠️ Skip image error: {e}")

    # ==========================
    # LOAD DATAFRAME
    # ==========================
    df = pd.read_excel(excel_path, sheet_name=real_name, header=None)

    # detectar header real
    for i in range(15):
        row = df.iloc[i].astype(str).str.lower()
        if "total cost usd" in row.values:
            df.columns = df.iloc[i]
            df = df[i+1:]
            break

    df = df.dropna(how="all")
    df = df.fillna("")

    # ==========================
    # MAP IMAGES
    # ==========================
    df["ImagePath"] = ""

    for idx in df.index:
        excel_row = idx + 2
        if excel_row in image_map:
            df.at[idx, "ImagePath"] = image_map[excel_row]

    return df

# ==========================
# PROCESS ALL
# ==========================
assy = process_sheet("ASSY", "assy")
qa = process_sheet("qa", "qa")
insertion = process_sheet("insertion", "insertion")
automation = process_sheet("automation", "automation")
fixtures = process_sheet("fixtures", "fixtures")

# ==========================
# SAVE OUTPUT
# ==========================
output_excel = "shopping_with_images.xlsx"

with pd.ExcelWriter(output_excel) as writer:
    if not assy.empty:
        assy.to_excel(writer, sheet_name="ASSY", index=False)
    if not qa.empty:
        qa.to_excel(writer, sheet_name="QA TOOLING", index=False)
    if not insertion.empty:
        insertion.to_excel(writer, sheet_name="INSERTION", index=False)
    if not automation.empty:
        automation.to_excel(writer, sheet_name="AUTOMATION", index=False)
    if not fixtures.empty:
        fixtures.to_excel(writer, sheet_name="FIXTURES & GAUGES", index=False)

print("✅ DONE → Generated:", output_excel)
